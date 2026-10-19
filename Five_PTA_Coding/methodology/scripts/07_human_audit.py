"""Human audit of the final PTA 2-5 codes (workflow: "have a person code a stratified sample (about 300 pairs across
steps and routes) to report accuracy of the final AI codes").

build: draws 300 pairs, stratified by how the final code is produced, excluding the 788 pairs of the AI pair sample:
       AI units (by input section), textile rules (by rule), ML keep1 / turn0, ML-rule conflicts and exact copies.
       Writes human_audit/human_audit_pairs.csv (blind: descriptions only) and _work/coded/human_audit_key.csv.
score: compares the filled file with final/master_coded_final.csv; accuracy by stratum and weighted totals
       (weights = stratum population / sample size) -> final/human_audit_report.json.

check: format check of an AI auditor file (see human_audit/AI_Audit_Prompt.md).
compare: AI auditors (audit_<coder>.csv) vs the final codes and vs each other; writes human_audit/audit_review.xlsx with
       the disagreements first for August's review, and final/ai_audit_report.json.

Run: /usr/local/bin/python3 _work/07_human_audit.py build | score | check <file> | compare <file> [<file> ...]
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from textile import load_texts  # noqa: E402

CODED = HERE / 'coded'
OUT = HERE.parent / 'human_audit'
FINAL = HERE.parent / 'final'
SEED = 20261019
ALLOC = {'ai_units': 120, 'rule_textile': 80, 'ml:keep1': 25, 'ml:turn0': 25, 'conflict': 20, 'exact_copy': 30}
GROUPS = [(1, 5, '01-05 animal'), (6, 15, '06-15 plant/fats'), (16, 24, '16-24 food'), (25, 27, '25-27 mineral'),
          (28, 38, '28-38 chemical'), (39, 40, '39-40 plastic/rubber'), (41, 49, '41-49 leather/wood/paper'),
          (50, 67, '50-67 textile/footwear'), (68, 71, '68-71 stone/glass/jewel'), (72, 83, '72-83 metals'),
          (84, 85, '84-85 machinery'), (86, 97, '86-97 transport/other')]


def group(ch):
    for a, b, g in GROUPS:
        if a <= ch <= b:
            return g
    return 'other'


def sqrt_alloc(sizes, total, floor):
    s = pd.Series(sizes, dtype=float)
    a = np.maximum(floor, np.floor(total * np.sqrt(s) / np.sqrt(s).sum())).astype(int)
    a = np.minimum(a, s.astype(int))
    while a.sum() > total:
        a[a.idxmax()] -= 1
    return a


def build():
    rng = np.random.default_rng(SEED)
    df = pd.read_csv(CODED / 'master_coded_stage2.csv', dtype=str, keep_default_na=False)
    key = pd.read_csv(CODED / 'sample_key.csv', dtype=str, keep_default_na=False)
    df = df[~df.master_id.isin(key.master_id) & (df.source != 'rule')]          # self-pairs are locked by rule
    df['stratum'] = ''
    df.loc[df.queue.isin(['unseen_other', 'ml_middle']), 'stratum'] = \
        'ai_units:' + df.input_hs6.str[:2].astype(int).map(group)
    df.loc[df.source == 'rule_textile', 'stratum'] = 'rule_textile:' + df.rule_id
    df.loc[df.source == 'ml', 'stratum'] = 'ml:' + df.band
    df.loc[df.queue == 'ml_rule_conflict', 'stratum'] = 'conflict:' + df.rule_id
    df.loc[df.source == 'exact_copy', 'stratum'] = 'exact_copy'
    df = df[df.stratum != '']
    picks = []
    for fam, n in ALLOC.items():
        sub = df[df.stratum.str.startswith(fam)]
        sizes = sub.stratum.value_counts()
        alloc = sqrt_alloc(sizes, n, 3 if fam in ('ai_units', 'rule_textile') else n) if len(sizes) > 1 else \
            pd.Series({sizes.index[0]: n})
        for s, k in alloc.items():
            g = sub[sub.stratum == s]
            take = g.iloc[rng.choice(len(g), size=int(k), replace=False)].copy()
            take['weight'] = len(g) / int(k)
            picks.append(take)
    P = pd.concat(picks).sample(frac=1, random_state=SEED).reset_index(drop=True)
    P['audit_id'] = [f'H{i + 1:03d}' for i in range(len(P))]
    T = load_texts()

    def desc(c):
        return f'{T.get(c[:4], "")} | {T.get(c, "")}'

    OUT.mkdir(exist_ok=True)
    blind = pd.DataFrame({'audit_id': P.audit_id, 'input_hs6': P.input_hs6, 'input_description': P.input_hs6.map(desc),
                          'output_hs6': P.output_hs6, 'output_description': P.output_hs6.map(desc),
                          'dv': '', 'route': '', 'step': '', 'reason': ''})
    blind.to_csv(OUT / 'human_audit_pairs.csv', index=False)
    try:
        blind.to_excel(OUT / 'human_audit_pairs.xlsx', index=False)
    except ImportError:
        pass
    P[['audit_id', 'master_id', 'stratum', 'weight']].to_csv(CODED / 'human_audit_key.csv', index=False)
    print(P.stratum.str.split(':').str[0].value_counts().to_dict(), len(P))


def score():
    key = pd.read_csv(CODED / 'human_audit_key.csv', dtype=str)
    f = OUT / 'human_audit_pairs.xlsx'
    h = pd.read_excel(f, dtype=str) if f.exists() else pd.read_csv(OUT / 'human_audit_pairs.csv', dtype=str)
    h = h[h.dv.isin(['0', '1'])]
    F = pd.read_csv(FINAL / 'master_coded_final.csv', dtype=str, keep_default_na=False).set_index('master_id')
    d = key.merge(h[['audit_id', 'dv']], on='audit_id')
    d['weight'] = d.weight.astype(float)
    d['human'] = d.dv.astype(int)
    d['final'] = d.master_id.map(F.DV).astype(int)
    d['strict'] = d.master_id.map(F.DV_strict).astype(int)
    d['source'] = d.master_id.map(F.source)
    d['ok'] = d.human == d.final
    fam = d.stratum.str.split(':').str[0]
    rep = {'coded': len(d), 'of': len(key),
           'accuracy_weighted_all': round(float(np.average(d.ok, weights=d.weight)), 4),
           'accuracy_unweighted': round(float(d.ok.mean()), 4),
           'DV_strict_accuracy_weighted': round(float(np.average(d.human == d.strict, weights=d.weight)), 4),
           'by_family': {k: {'n': len(g), 'acc': round(float(np.average(g.ok, weights=g.weight)), 4),
                             'human1_final0': int(((g.human == 1) & (g.final == 0)).sum()),
                             'human0_final1': int(((g.human == 0) & (g.final == 1)).sum())}
                         for k, g in d.groupby(fam)},
           'by_stratum': {k: {'n': len(g), 'acc': round(float(g.ok.mean()), 4)} for k, g in d.groupby('stratum')}}
    json.dump(rep, open(FINAL / 'human_audit_report.json', 'w'), indent=1)
    d.to_csv(FINAL / 'human_audit_scored.csv', index=False)
    print(json.dumps(rep, indent=1))


ROUTES1 = {'same_good', 'forward', 'ingredient', 'component', 'reconstitute', 'mixture', 'recycle_waste',
           'recycle_direct', 'feed', 'sibling_intermediate'}


def check(path):
    """format check of an AI auditor file (audit_id,dv,route,step,reason; one row per item, input order)"""
    import csv
    path = Path(path) if Path(path).is_absolute() else HERE.parent / path
    rows = list(csv.reader(open(path, newline='', encoding='utf-8')))
    err = []
    if rows[0] != ['audit_id', 'dv', 'route', 'step', 'reason']:
        err.append(f'header {rows[0]}')
    bad = [k + 1 for k, r in enumerate(rows) if r and len(r) != 5]
    if bad:
        err.append(f'{len(bad)} line(s) with a wrong number of fields (unquoted comma?), e.g. {bad[:10]}')
    if not err:
        a = pd.read_csv(path, dtype=str, keep_default_na=False)
        ids = pd.read_csv(OUT / 'human_audit_pairs.csv', dtype=str).audit_id
        if len(a) != len(ids) or (a.audit_id.values != ids.values).any():
            err.append(f'items differ from the input (rows {len(a)} vs {len(ids)}, or order)')
        for i, r in a.iterrows():
            b = []
            if r.dv not in ('0', '1'):
                b.append('dv')
            elif r.dv == '1' and r.route not in ROUTES1:
                b.append('route for dv 1')
            elif r.dv == '0' and r.route != 'none':
                b.append('route for dv 0 must be none')
            if r.route == 'same_good' and r.step != '2':
                b.append('same_good needs step 2')
            if r.dv == '0' and r.step not in ('3', '4', '5', '6'):
                b.append('dv 0 needs step 3-6')
            if r.dv == '1' and r.route != 'same_good' and r.step != '':
                b.append('step must be empty')
            if len(r.reason.strip()) < 15:
                b.append('reason too short')
            if b:
                err.append(f'row {i + 2} {r.audit_id}: {", ".join(b)}')
    print(f'{path.name}: {"OK" if not err else f"{len(err)} problem(s)"}')
    for e in err[:25]:
        print('   ', e)
    sys.exit(0 if not err else 1)


def kappa(a, b):
    a, b = np.asarray(a, int), np.asarray(b, int)
    po = (a == b).mean()
    pe = a.mean() * b.mean() + (1 - a.mean()) * (1 - b.mean())
    return round(float((po - pe) / (1 - pe)), 4) if pe < 1 else 1.0


def compare(paths):
    """AI auditors vs the final codes: accuracy by stratum family, agreement between auditors, and a review workbook
    for August with the disagreements first (human_audit/audit_review.xlsx; final/ai_audit_report.json)"""
    key = pd.read_csv(CODED / 'human_audit_key.csv', dtype=str)
    pairs = pd.read_csv(OUT / 'human_audit_pairs.csv', dtype=str, keep_default_na=False)
    F = pd.read_csv(FINAL / 'master_coded_final.csv', dtype=str, keep_default_na=False).set_index('master_id')
    d = key.merge(pairs[['audit_id', 'input_hs6', 'input_description', 'output_hs6', 'output_description']],
                  on='audit_id')
    d['weight'] = d.weight.astype(float)
    d['family'] = d.stratum.str.split(':').str[0]
    for c in ['DV', 'route', 'source', 'uncertain', 'uncertain_reason']:
        d[f'final_{c}'] = d.master_id.map(F[c])
    d['final_DV'] = d.final_DV.astype(int)
    names = []
    for path in paths:
        path = Path(path) if Path(path).is_absolute() else HERE.parent / path
        nm = path.stem.replace('audit_', '')
        a = pd.read_csv(path, dtype=str, keep_default_na=False).set_index('audit_id')
        d[f'{nm}_dv'] = d.audit_id.map(a.dv).astype(int)
        d[f'{nm}_route'] = d.audit_id.map(a.route)
        d[f'{nm}_reason'] = d.audit_id.map(a.reason)
        names.append(nm)
    rep = {'auditors': names, 'items': len(d)}
    for nm in names:
        ok = d[f'{nm}_dv'] == d.final_DV
        rep[nm] = {'agree_with_final_weighted': round(float(np.average(ok, weights=d.weight)), 4),
                   'agree_with_final_unweighted': round(float(ok.mean()), 4),
                   'auditor1_final0': int(((d[f'{nm}_dv'] == 1) & (d.final_DV == 0)).sum()),
                   'auditor0_final1': int(((d[f'{nm}_dv'] == 0) & (d.final_DV == 1)).sum()),
                   'by_family': {k: {'n': len(g), 'agree': round(float(np.average(g[f'{nm}_dv'] == g.final_DV,
                                                                                   weights=g.weight)), 4)}
                                 for k, g in d.groupby('family')}}
    av = d[[f'{nm}_dv' for nm in names]]
    same = av.nunique(axis=1) == 1
    if len(names) > 1:
        rep['between_auditors'] = {'agree': round(float(same.mean()), 4),
                                   'kappa': kappa(av.iloc[:, 0], av.iloc[:, 1])}
        cons = av.iloc[:, 0]
        okc = (cons == d.final_DV)[same]
        rep['final_vs_auditor_consensus'] = {
            'items_where_auditors_agree': int(same.sum()),
            'final_agrees_weighted': round(float(np.average(okc, weights=d.weight[same])), 4),
            'by_family': {k: {'n': int(m.sum()), 'agree': round(float((cons == d.final_DV)[m].mean()), 4)}
                          for k in sorted(d.family.unique()) for m in [same & (d.family == k)] if m.any()}}
    allok = same & (av.iloc[:, 0] == d.final_DV)
    d['status'] = np.where(allok, '3 all agree',
                           np.where(same, '1 auditors agree, final differs', '2 auditors disagree'))
    rep['review_counts'] = d.status.value_counts().sort_index().to_dict()
    d = d.sort_values(['status', 'family', 'audit_id'])
    d['August_check'] = ''
    d['August_note'] = ''
    cols = (['status', 'audit_id', 'family', 'input_hs6', 'input_description', 'output_hs6', 'output_description',
             'final_DV', 'final_route', 'final_source', 'final_uncertain']
            + [f'{nm}_{c}' for nm in names for c in ('dv', 'route', 'reason')] + ['August_check', 'August_note'])
    out = OUT / 'audit_review.xlsx'
    d[cols].to_excel(out, index=False)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        ws.freeze_panes = 'D2'
        for i, w in enumerate([24, 8, 11, 9, 40, 9, 40, 8, 18, 14, 8] + [7, 16, 40] * len(names) + [16, 30], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        fills = {'1': 'F8D7DA', '2': 'FFF3CD', '3': 'FFFFFF'}
        for row in ws.iter_rows(min_row=2):
            f = PatternFill('solid', fgColor=fills[str(row[0].value)[0]])
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
                c.fill = f
        for c in ws[1]:
            c.font = Font(bold=True)
        dv = DataValidation(type='list', formula1='"final correct,final wrong,unsure"', allow_blank=True)
        ws.add_data_validation(dv)
        col = openpyxl.utils.get_column_letter(len(cols) - 1)
        dv.add(f'{col}2:{col}{ws.max_row}')
        wb.save(out)
    except ImportError:
        pass
    json.dump(rep, open(FINAL / 'ai_audit_report.json', 'w'), indent=1)
    print(json.dumps(rep, indent=1))


if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'check':
        check(sys.argv[2])
    elif cmd == 'compare':
        compare(sys.argv[2:])
    else:
        {'build': build, 'score': score}[cmd]()
